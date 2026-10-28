from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "output" / "report"
REPORT_MD = REPORT_DIR / "REPORT.md"
DOCX_OUT = REPORT_DIR / "Final_Report.docx"
PDF_OUT = REPORT_DIR / "Final_Report.pdf"


def parse_report_md(md_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    sections = []
    current = {"title": "", "body": []}

    for line in lines:
        if line.startswith("# "):
            if current["title"] or current["body"]:
                sections.append(current)
            current = {"title": line[2:].strip(), "body": []}
        elif line.startswith("## "):
            if current["title"] or current["body"]:
                sections.append(current)
            current = {"title": line[3:].strip(), "body": []}
        else:
            current["body"].append(line)

    if current["title"] or current["body"]:
        sections.append(current)

    return sections


def export_docx(sections):
    doc = Document()
    doc.add_heading("Prediction-Market Final Report", level=0)

    for sec in sections:
        if sec["title"]:
            level = 1 if sec["title"].lower().startswith("macro prediction") else 2
            doc.add_heading(sec["title"], level=level)

        for line in sec["body"]:
            line = line.strip()
            if not line:
                continue

            m = re.match(r"!\[\]\((.*?)\)", line)
            if m:
                img_name = m.group(1)
                img_path = REPORT_DIR / img_name
                if img_path.exists():
                    doc.add_picture(str(img_path), width=Inches(6.2))
                continue

            if line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)

    doc.save(DOCX_OUT)


def _draw_wrapped_text(c: canvas.Canvas, text: str, x: float, y: float, max_width: float, leading: float = 14):
    words = text.split()
    line = ""
    for w in words:
        test = (line + " " + w).strip()
        if c.stringWidth(test, "Helvetica", 11) <= max_width:
            line = test
        else:
            c.drawString(x, y, line)
            y -= leading
            line = w
    if line:
        c.drawString(x, y, line)
        y -= leading
    return y


def export_pdf(sections):
    c = canvas.Canvas(str(PDF_OUT), pagesize=A4)
    page_w, page_h = A4

    margin = 40
    y = page_h - margin
    max_w = page_w - 2 * margin

    c.setFont("Helvetica-Bold", 18)
    c.drawString(margin, y, "Prediction-Market Final Report")
    y -= 30

    for sec in sections:
        if y < 100:
            c.showPage()
            y = page_h - margin

        if sec["title"]:
            c.setFont("Helvetica-Bold", 14)
            c.drawString(margin, y, sec["title"])
            y -= 20

        c.setFont("Helvetica", 11)
        for line in sec["body"]:
            line = line.strip()
            if not line:
                continue

            m = re.match(r"!\[\]\((.*?)\)", line)
            if m:
                img_name = m.group(1)
                img_path = REPORT_DIR / img_name
                if img_path.exists():
                    with Image.open(img_path) as im:
                        iw, ih = im.size
                    scale = min(max_w / iw, 300 / ih)
                    dw, dh = iw * scale, ih * scale
                    if y - dh < 50:
                        c.showPage()
                        y = page_h - margin
                    c.drawImage(ImageReader(str(img_path)), margin, y - dh, width=dw, height=dh)
                    y -= dh + 14
                continue

            text = line[2:] if line.startswith("- ") else line
            if y < 60:
                c.showPage()
                y = page_h - margin
                c.setFont("Helvetica", 11)
            y = _draw_wrapped_text(c, text, margin, y, max_w)

        y -= 8

    c.save()


def main():
    if not REPORT_MD.exists():
        raise FileNotFoundError(f"Missing report markdown: {REPORT_MD}")

    sections = parse_report_md(REPORT_MD)
    export_docx(sections)
    export_pdf(sections)

    print(f"[export] DOCX -> {DOCX_OUT}")
    print(f"[export] PDF  -> {PDF_OUT}")


if __name__ == "__main__":
    main()
